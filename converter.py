import os
import shutil
import subprocess
import tempfile
from pathlib import Path

# Safe lazy imports for Termux / limited Linux environments
try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    Image = None
    HAS_PIL = False

try:
    from pdf2docx import Converter
    HAS_PDF2DOCX = True
except ImportError:
    Converter = None
    HAS_PDF2DOCX = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    docx = None
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    Presentation = None
    HAS_PPTX = False

try:
    import fitz  # PyMuPDF
    HAS_FITZ = True
except ImportError:
    fitz = None
    HAS_FITZ = False

try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    pypdf = None
    HAS_PYPDF = False

SUPPORTED_EXTENSIONS = {
    ".pdf": ["docx", "xlsx", "pptx", "images"],
    ".docx": ["pdf"],
    ".xlsx": ["pdf"],
    ".pptx": ["pdf"],
    ".jpg": ["pdf"],
    ".jpeg": ["pdf"],
    ".png": ["pdf"],
}

LIBREOFFICE_PATHS = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    shutil.which("soffice") or "",
    shutil.which("libreoffice") or "",
]


def get_libreoffice_bin() -> str:
    for p in LIBREOFFICE_PATHS:
        if p and os.path.exists(p):
            return p
    return ""


def convert_pdf_to_docx(input_path: str, output_path: str) -> str:
    if not HAS_PDF2DOCX:
        # Fallback using pypdf + docx if pdf2docx is not installed
        if HAS_PYPDF and HAS_DOCX:
            reader = pypdf.PdfReader(input_path)
            doc = docx.Document()
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    for line in text.split("\n"):
                        doc.add_paragraph(line)
                    doc.add_page_break()
            doc.save(output_path)
            return output_path
        raise RuntimeError("Modul konversi PDF ke DOCX (pdf2docx/pypdf) belum terpasang.")
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()
    return output_path


def convert_pdf_to_xlsx(input_path: str, output_path: str) -> str:
    if not HAS_OPENPYXL:
        raise RuntimeError("Modul openpyxl belum terpasang.")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted"
    current_row = 1

    if HAS_FITZ:
        doc = fitz.open(input_path)
        for page in doc:
            tabs = getattr(page, "find_tables", lambda: None)()
            if tabs and getattr(tabs, "tables", None):
                for table in tabs.tables:
                    for row in table.extract():
                        for col_idx, val in enumerate(row, 1):
                            ws.cell(row=current_row, column=col_idx, value=val)
                        current_row += 1
                    current_row += 1
            else:
                lines = page.get_text("text").splitlines()
                for line in lines:
                    if line.strip():
                        ws.cell(row=current_row, column=1, value=line.strip())
                        current_row += 1
                current_row += 1
    elif HAS_PYPDF:
        reader = pypdf.PdfReader(input_path)
        for page in reader.pages:
            text = page.extract_text()
            if text:
                for line in text.splitlines():
                    if line.strip():
                        ws.cell(row=current_row, column=1, value=line.strip())
                        current_row += 1
                current_row += 1
    else:
        raise RuntimeError("Modul ekstraksi PDF (pypdf/fitz) belum terpasang.")

    wb.save(output_path)
    return output_path


def convert_pdf_to_pptx(input_path: str, output_path: str) -> str:
    if not HAS_PPTX or not HAS_FITZ:
        raise RuntimeError("Modul konversi PDF ke PPTX (python-pptx/fitz) belum terpasang.")
    prs = Presentation()
    prs.slide_width = fitz.paper_size("a4")[1] * 72 * 10
    doc = fitz.open(input_path)
    blank_slide_layout = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmp_dir:
        for idx, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            img_path = os.path.join(tmp_dir, f"page_{idx}.png")
            pix.save(img_path)
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width)
    prs.save(output_path)
    return output_path


def convert_pdf_to_images(input_path: str, output_dir: str) -> list[str]:
    if not HAS_FITZ:
        raise RuntimeError("Ekstraksi gambar dari PDF membutuhkan PyMuPDF.")
    doc = fitz.open(input_path)
    image_paths = []
    for idx, page in enumerate(doc):
        pix = page.get_pixmap(dpi=200)
        img_path = os.path.join(output_dir, f"page_{idx + 1}.png")
        pix.save(img_path)
        image_paths.append(img_path)
    return image_paths


def convert_image_to_pdf(input_path: str, output_path: str) -> str:
    if not HAS_PIL:
        raise RuntimeError("Pillow belum terpasang untuk konversi gambar.")
    img = Image.open(input_path)
    if img.mode != "RGB":
        img = img.convert("RGB")
    img.save(output_path, "PDF", resolution=100.0)
    return output_path


def convert_office_to_pdf(input_path: str, output_dir: str) -> str:
    lo_bin = get_libreoffice_bin()
    if lo_bin:
        cmd = [lo_bin, "--headless", "--convert-to", "pdf", "--outdir", output_dir, input_path]
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
        base_name = Path(input_path).stem + ".pdf"
        out_file = os.path.join(output_dir, base_name)
        if os.path.exists(out_file):
            return out_file

    ext = Path(input_path).suffix.lower()
    if ext == ".docx" and HAS_DOCX:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        out_file = os.path.join(output_dir, Path(input_path).stem + ".pdf")
        doc = docx.Document(input_path)
        c = canvas.Canvas(out_file, pagesize=letter)
        y = 750
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                c.drawString(50, y, text[:90])
                y -= 20
                if y < 50:
                    c.showPage()
                    y = 750
        c.save()
        return out_file
    elif ext == ".xlsx" and HAS_OPENPYXL:
        wb = openpyxl.load_workbook(input_path, data_only=True)
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        out_file = os.path.join(output_dir, Path(input_path).stem + ".pdf")
        c = canvas.Canvas(out_file, pagesize=letter)
        y = 750
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            c.drawString(50, y, f"--- Sheet: {sheet} ---")
            y -= 25
            for row in ws.iter_rows(values_only=True):
                line = " | ".join(str(v) for v in row if v is not None)
                if line.strip():
                    c.drawString(50, y, line[:85])
                    y -= 18
                    if y < 50:
                        c.showPage()
                        y = 750
        c.save()
        return out_file
    else:
        raise RuntimeError("LibreOffice atau engine PDF tidak ditemukan untuk format ini.")
