from datetime import datetime, timezone
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from voice_stats import VoiceStatsStore


def format_duration(seconds: int) -> str:
    hours, remainder = divmod(max(0, int(seconds)), 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours}h {minutes}m {seconds}s"


def _format_time(value: datetime, store: VoiceStatsStore) -> str:
    return value.astimezone(store.timezone).strftime("%Y-%m-%d %H:%M:%S")


HEADER_FILL = PatternFill("solid", fgColor="FF234E70")
ALT_FILL = PatternFill("solid", fgColor="FFEAF1F8")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _setup_sheet(sheet, title: str, headers: list[str], widths: list[int]) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(name="Calibri", size=16, bold=True)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[3].height = 36
    for column, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(3, column, header)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = CENTER
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "C4"
    sheet.sheet_view.showGridLines = True


def _add_row(sheet, values: list[object]) -> None:
    row_number = sheet.max_row + 1
    fill = ALT_FILL if row_number % 2 == 0 else WHITE_FILL
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column, value)
        cell.font = Font(name="Calibri", size=11)
        cell.fill = fill
        cell.border = THIN_BORDER
        cell.alignment = CENTER


def _finish_sheet(sheet, column_count: int) -> None:
    sheet.auto_filter.ref = f"A3:{sheet.cell(sheet.max_row, column_count).coordinate}"


def export_period(store: VoiceStatsStore, start: datetime, end: datetime, output_path: str | Path) -> Path:
    summary = store.summarize(start, end)
    workbook = Workbook()
    sessions_sheet = workbook.active
    sessions_sheet.title = "Sessions"
    _setup_sheet(
        sessions_sheet,
        "Voice Activity Sessions",
        ["Member", "User ID", "Voice Channel", "Join Time", "Leave Time", "Duration"],
        [28, 22, 28, 23, 23, 18],
    )
    now = datetime.now(timezone.utc)
    clipped_rows = []
    for row in summary["rows"]:
        row_start = max(store._parse(row["joined_at"]), start)
        row_end = min(store._parse(row["left_at"]) if row["left_at"] else now, end)
        if row_end <= row_start:
            continue
        clipped_rows.append((row, row_start, row_end))
        _add_row(sessions_sheet, [
            row["username"], row["user_id"], row["channel_name"], _format_time(row_start, store),
            _format_time(row_end, store), format_duration(int((row_end - row_start).total_seconds())),
        ])
    _finish_sheet(sessions_sheet, 6)

    daily_sheet = workbook.create_sheet("Daily Summary")
    _setup_sheet(
        daily_sheet,
        "Voice Activity Daily Summary",
        ["Date", "Total Voice Time", "Total Sessions", "Active Members"],
        [18, 24, 20, 20],
    )
    for date, value in sorted(summary["daily"].items()):
        _add_row(daily_sheet, [date, format_duration(value["seconds"]), len(value["sessions"]), len(value["members"])])
    _finish_sheet(daily_sheet, 4)

    member_sheet = workbook.create_sheet("Member Summary")
    _setup_sheet(
        member_sheet,
        "Voice Activity Member Summary",
        ["Member", "User ID", "Total Voice Time", "Number of Sessions", "Average Session Duration"],
        [28, 22, 24, 20, 28],
    )
    member_totals: dict[int, dict] = {}
    for row, row_start, row_end in clipped_rows:
        total = member_totals.setdefault(row["user_id"], {"name": row["username"], "seconds": 0, "sessions": 0})
        total["name"] = row["username"]
        total["seconds"] += int((row_end - row_start).total_seconds())
        total["sessions"] += 1
    for user_id, total in sorted(member_totals.items(), key=lambda item: item[1]["seconds"], reverse=True):
        _add_row(member_sheet, [
            total["name"], user_id, format_duration(total["seconds"]), total["sessions"],
            format_duration(total["seconds"] // total["sessions"]),
        ])
    _finish_sheet(member_sheet, 5)

    channel_sheet = workbook.create_sheet("Channel Summary")
    _setup_sheet(
        channel_sheet,
        "Voice Activity Channel Summary",
        ["Voice Channel", "Total Voice Time", "Number of Sessions"],
        [28, 24, 20],
    )
    channel_totals: dict[int, dict] = {}
    for row, row_start, row_end in clipped_rows:
        total = channel_totals.setdefault(row["channel_id"], {"name": row["channel_name"], "seconds": 0, "sessions": 0})
        total["name"] = row["channel_name"]
        total["seconds"] += int((row_end - row_start).total_seconds())
        total["sessions"] += 1
    for total in sorted(channel_totals.values(), key=lambda item: item["seconds"], reverse=True):
        _add_row(channel_sheet, [total["name"], format_duration(total["seconds"]), total["sessions"]])
    _finish_sheet(channel_sheet, 3)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Write beside the destination so a failed save cannot truncate a previous export.
    descriptor, temporary_name = tempfile.mkstemp(prefix=".voice-export-", suffix=".xlsx", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return output
