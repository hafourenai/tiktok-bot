"""Generate reference fixtures with the ORIGINAL third-party libraries.

Run once; the committed binaries in tests/fixtures/ are the ground truth the
stdlib reader is validated against. Regenerating them re-introduces the
dependency, so this script is deliberately not part of the test suite.

    venv\\Scripts\\python.exe tests\\make_fixtures.py
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(HERE, "fixtures")
os.makedirs(FIXTURES, exist_ok=True)


def out(name):
    return os.path.join(FIXTURES, name)


def make_text_pdf():
    from reportlab.pdfgen import canvas

    path = out("text_reportlab.pdf")
    pdf = canvas.Canvas(path, pagesize=(300, 500))
    pdf.drawString(20, 450, "Invoice 123")
    pdf.drawString(20, 430, "Total: Rp 1.500.000")
    pdf.showPage()
    pdf.setPageSize((500, 300))
    pdf.drawString(20, 250, "Second page")
    pdf.save()
    return path


def make_table_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

    path = out("table_reportlab.pdf")
    body = Table(
        [["Product", "Qty", "Amount"],
         ["Kopi", "2", "30.000"],
         ["Teh", "1", "5.000"]],
        colWidths=[120, 40, 80],
    )
    body.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    footer = Table([["Total", "", "35.000"]], colWidths=[120, 40, 80])
    footer.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("SPAN", (0, 0), (1, 0)),
    ]))
    story = [body, Spacer(1, 12), footer]
    SimpleDocTemplate(path, pagesize=letter).build(story)
    return path


def make_unruled_pdf():
    from reportlab.pdfgen import canvas

    path = out("unruled_reportlab.pdf")
    pdf = canvas.Canvas(path, pagesize=(400, 300))
    rows = [("Nama", "Qty", "Harga"),
            ("Beras", "5", "65000"),
            ("Minyak", "2", "32000"),
            ("Gula", "1", "15000")]
    y = 250
    for row in rows:
        x = 30
        for value in row:
            pdf.drawString(x, y, value)
            x += 120
        y -= 20
    pdf.save()
    return path


def make_multipage_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = out("multipage_reportlab.pdf")
    pdf = canvas.Canvas(path, pagesize=letter)
    for index in range(5):
        pdf.drawString(72, 720, f"Halaman nomor {index + 1}")
        if index < 4:
            pdf.setPageSize((400, 500))
            pdf.showPage()
            pdf.setPageSize(letter)
    pdf.save()
    return path


def make_winansi_pdf():
    from reportlab.pdfgen import canvas

    path = out("winansi_reportlab.pdf")
    pdf = canvas.Canvas(path)
    pdf.setFont("Helvetica", 14)
    pdf.drawString(60, 700, "Harga: Rp 1.500.000 (net)")
    pdf.drawString(60, 680, "Nama replicate: cafe - stra\u00dfe")
    pdf.setFont("Times-Roman", 12)
    pdf.drawString(60, 660, "Total 45,60 kg \u2014 selesai")
    pdf.showPage()
    pdf.setFont("Courier", 12)
    pdf.drawString(60, 700, "MONOSPACE 0123456789 invoice")
    pdf.save()
    return path


def make_scan_pdf():
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    import struct
    import zlib

    # 2x2 raw RGB bitmap wrapped in a PNG, then embedded as an image XObject.
    width = height = 2
    raw = b"".join(b"\x00" + bytes([255, 255, 255]) * width for _ in range(height))

    def chunk(tag, payload):
        body = tag + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(
            ">I", zlib.crc32(body) & 0xFFFFFFFF)

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw))
           + chunk(b"IEND", b""))

    path = out("scan_reportlab.pdf")
    pdf = canvas.Canvas(path, pagesize=(200, 200))
    pdf.drawImage(ImageReader(__import__("io").BytesIO(png)), 20, 20, 160, 160)
    pdf.save()
    return path


def make_docx():
    import docx

    path = out("sample_docx.docx")
    document = docx.Document()
    document.add_paragraph("Laporan keuangan")
    document.add_paragraph("Baris kedua untuk pengujian")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Jumlah"
    table.cell(1, 0).text = "Kopi"
    table.cell(1, 1).text = "2"
    document.save(path)
    return path


def make_xlsx():
    import openpyxl

    path = out("sample_xlsx.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Nama", "Qty", "Harga"])
    ws.append(["Beras", 5, 65000])
    ws.append(["Minyak", 2, 32000])
    wb.save(path)
    wb.close()
    return path


def make_pptx():
    import pptx
    from pptx.util import Emu

    path = out("sample_pptx.pptx")
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Emu(500000), Emu(500000),
                             Emu(4000000), Emu(1000000)).text = "Slide pertama"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide2.shapes.add_textbox(Emu(500000), Emu(500000),
                                    Emu(4000000), Emu(1000000))
    box.text_frame.text = "Slide kedua"
    prs.save(path)
    return path


BUILDERS = [
    make_text_pdf, make_table_pdf, make_unruled_pdf, make_multipage_pdf,
    make_winansi_pdf, make_scan_pdf, make_docx, make_xlsx, make_pptx,
]


def main():
    failures = []
    for builder in BUILDERS:
        try:
            path = builder()
            print(f"  OK  {os.path.basename(path)}  ({os.path.getsize(path)} bytes)")
        except Exception as exc:  # noqa: BLE001 - report and continue
            failures.append((builder.__name__, exc))
            print(f"  ERR {builder.__name__}: {exc}")
    if failures:
        print(f"\n{len(failures)} fixture gagal dibuat.")
        return 1
    print(f"\nSemua fixture dibuat di {FIXTURES}")
    return 0


if __name__ == "__main__":
    sys.exit(main())