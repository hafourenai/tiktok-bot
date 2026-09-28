from datetime import datetime, timezone

from openpyxl import load_workbook

from voice_export import export_period
from voice_stats import VoiceStatsStore


UTC = timezone.utc


def timestamp(year, month, day, hour=0, minute=0):
    return datetime(year, month, day, hour, minute, tzinfo=UTC)


def test_move_creates_two_sessions(tmp_path):
    store = VoiceStatsStore(tmp_path / "voice.db")
    store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 9, 1, 9))
    store.close_session(1, 10, timestamp(2026, 9, 1, 10))
    store.start_session(1, 10, "Aldi", 200, "Nongkrong", timestamp(2026, 9, 1, 10))
    store.close_session(1, 10, timestamp(2026, 9, 1, 11, 30))

    summary = store.summarize(timestamp(2026, 9, 1, 0), timestamp(2026, 9, 2, 0))
    assert summary["sessions"] == 2
    assert summary["total_seconds"] == 150 * 60
    assert summary["channels"] == {"Nongkrong": 90 * 60, "Gaming": 60 * 60}


def test_active_session_is_not_duplicated_after_recovery(tmp_path):
    store = VoiceStatsStore(tmp_path / "voice.db")
    assert store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 9, 1, 9))
    assert not store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 9, 1, 10))
    assert len(store.active_sessions(1)) == 1


def test_midnight_duration_is_split_by_local_date(tmp_path):
    store = VoiceStatsStore(tmp_path / "voice.db", "Asia/Jakarta")
    store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 9, 1, 16, 30))
    store.close_session(1, 10, timestamp(2026, 9, 1, 18, 30))

    summary = store.summarize(timestamp(2026, 9, 1, 0), timestamp(2026, 9, 2, 23))
    assert summary["daily"]["2026-09-01"]["seconds"] == 30 * 60
    assert summary["daily"]["2026-09-02"]["seconds"] == 90 * 60


def test_delete_month_preserves_session_portions_outside_period(tmp_path):
    store = VoiceStatsStore(tmp_path / "voice.db")
    store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 8, 31, 23))
    store.close_session(1, 10, timestamp(2026, 10, 1, 1))
    start, end = timestamp(2026, 9, 1), timestamp(2026, 10, 1)

    assert store.delete_period(start, end) == 1
    august = store.summarize(timestamp(2026, 8, 31, 0), start)
    october = store.summarize(end, timestamp(2026, 10, 2))
    assert august["total_seconds"] == 60 * 60
    assert october["total_seconds"] == 60 * 60


def test_export_uses_template_style_and_keeps_same_named_members_separate(tmp_path):
    store = VoiceStatsStore(tmp_path / "voice.db")
    start, end = timestamp(2026, 9, 1), timestamp(2026, 10, 1)
    store.start_session(1, 10, "Aldi", 100, "Gaming", timestamp(2026, 9, 1, 9))
    store.close_session(1, 10, timestamp(2026, 9, 1, 10))
    store.start_session(1, 20, "Aldi", 200, "Nongkrong", timestamp(2026, 9, 2, 9))
    store.close_session(1, 20, timestamp(2026, 9, 2, 11))

    output = export_period(store, start, end, tmp_path / "voice_history_2026-09.xlsx")
    workbook = load_workbook(output)

    assert workbook.sheetnames == ["Sessions", "Daily Summary", "Member Summary", "Channel Summary"]
    sessions = workbook["Sessions"]
    assert any(str(cell_range) == "A1:F1" for cell_range in sessions.merged_cells.ranges)
    assert sessions["A1"].value == "Voice Activity Sessions"
    assert sessions["A3"].fill.fgColor.rgb == "FF234E70"
    assert sessions["A4"].fill.fgColor.rgb == "FFEAF1F8"
    assert sessions.freeze_panes == "C4"
    assert sessions.auto_filter.ref == "A3:F5"

    members = workbook["Member Summary"]
    assert members.max_row == 5
    assert {members["B4"].value, members["B5"].value} == {10, 20}
