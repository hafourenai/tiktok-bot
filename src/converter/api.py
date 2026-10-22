"""Document conversion API with OCR fallback for scanned pages.

Public API for converting between document formats.
Supports: PDF ↔ DOCX/XLSX/PPTX, PDF ↔ Images, Image → PDF, Office → PDF.
"""

import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from .ocr import OCRUnavailable, OCRError, image_to_text, is_ocr_available
from .pipeline import new_workspace, cleanup_stale_workspaces

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf": ["docx", "xlsx", "pptx", "images", "compress_light", "compress_medium", "compress_high"],
    ".docx": ["pdf"], ".xlsx": ["pdf"], ".pptx": ["pdf"],
    ".jpg": ["pdf"], ".jpeg": ["pdf"], ".png": ["pdf"],
}


def get_libreoffice_bin() -> str:
    """Locate LibreOffice binary for Office → PDF conversion."""
    candidates = [
        os.environ.get("LIBREOFFICE_BIN", ""),
        shutil.which("soffice"), shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    return next((p for p in candidates if p and Path(p).is_file()), "")


def get_poppler_bin() -> str:
    """Locate pdftoppm from poppler for PDF rasterization."""
    candidates = [
        os.environ.get("POPPLER_BIN", ""),
        shutil.which("pdftoppm"),
        r"C:\Program Files\poppler\Library\bin\pdftoppm.exe",
        r"C:\Program Files\poppler\bin\pdftoppm.exe",
        "/usr/bin/pdftoppm", "/usr/local/bin/pdftoppm",
        "/data/data/com.termux/files/usr/bin/pdftoppm",
    ]
    return next((p for p in candidates if p and Path(p).is_file()), "")


POPPLER_MISSING = (
    "Render PDF membutuhkan pdftoppm dari poppler.\n"
    "Windows: download dari oschwartz10612/poppler-windows, extract, tambah ke PATH.\n"
    "Termux: pkg install poppler.\n"
    "Ekstraksi teks tetap berjalan tanpa poppler."
)


def _render_pdf(input_path: str, output_dir: str, dpi: int = 150) -> list[str]:
    """Rasterize PDF pages to PNG using pdftoppm."""
    binary = get_poppler_bin()
    if not binary:
        raise RuntimeError(POPPLER_MISSING)
    
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = output_dir / "page"
    cmd = [binary, "-png", "-r", str(int(dpi)), str(input_path), str(prefix)]
    
    try:
        result = subprocess.run(cmd, capture_output=True, timeout=180)
    except (OSError, subprocess.SubprocessError) as exc:
        raise RuntimeError(f"Gagal menjalankan pdftoppm: {exc}") from exc
    
    if result.returncode != 0:
        detail = (result.stderr or b"").decode("utf-8", "replace").strip()
        raise RuntimeError(f"pdftoppm gagal (kode {result.returncode}): {detail}")
    
    produced = sorted(output_dir.glob("page-*.png"),
                     key=lambda p: _page_sort_key(p.stem))
    if not produced:
        raise RuntimeError("pdftoppm tidak menghasilkan gambar.")
    
    return [str(path) for path in produced]


def _page_sort_key(stem: str) -> tuple:
    """Sort page filenames numerically (page-1, page-2, etc)."""
    suffix = stem.rsplit("-", 1)[-1]
    return (0, int(suffix)) if suffix.isdigit() else (1, 0)


def _pdf_pages(input_path: str, workspace: str) -> list[tuple[str, list]]:
    """Extract text and tables from PDF pages with OCR fallback.
    
    Returns list of (text, tables) tuples per page.
    Raises RuntimeError if PDF cannot be read or has no extractable content.
    """
    from .lib.docconv.pdf import PdfReadError, extract_pages
    
    try:
        extracted = extract_pages(input_path)
    except PdfReadError as exc:
        raise RuntimeError(str(exc)) from exc
    
    pages = [(page.text, page.tables) for page in extracted]
    if not pages:
        raise RuntimeError("PDF tidak memiliki halaman.")
    
    # Detect scanned pages (no text + no tables)
    needs_ocr = [idx for idx, (text, tables) in enumerate(pages) 
                 if not text.strip() and not tables]
    
    if needs_ocr:
        if not is_ocr_available():
            logger.warning("Halaman scan terdeteksi tapi OCR tidak tersedia")
        else:
            _apply_ocr_fallback(input_path, pages, needs_ocr, workspace)
    
    # Final check: ensure all pages have content
    if any(not text.strip() and not tables for text, tables in pages):
        msg = (
            "PDF memiliki halaman tanpa teks yang dapat diekstrak (scan/gambar atau kosong). "
            "Pasang tesseract-ocr untuk OCR otomatis atau gunakan PPT/gambar."
        )
        raise RuntimeError(msg)
    
    return pages


def _apply_ocr_fallback(input_path: str, pages: list, 
                       needs_ocr: list[int], workspace: str) -> None:
    """Apply OCR to scanned pages, updating pages list in-place."""
    try:
        rendered_images = _render_pdf(input_path, workspace, dpi=200)
    except RuntimeError as e:
        logger.warning("Tidak dapat merender PDF untuk OCR: %s", e)
        return
    
    ocr_count = 0
    for idx in needs_ocr:
        if idx >= len(rendered_images):
            continue
        try:
            ocr_text = image_to_text(rendered_images[idx])
            if ocr_text:
                _, tables = pages[idx]
                pages[idx] = (ocr_text, tables)
                ocr_count += 1
                logger.info("OCR berhasil untuk halaman %d", idx + 1)
        except (OCRUnavailable, OCRError) as e:
            logger.warning("OCR gagal untuk halaman %d: %s", idx + 1, e)
            continue
    
    if ocr_count > 0:
        logger.info("OCR berhasil untuk %d halaman", ocr_count)


def convert_pdf_to_docx(input_path: str, output_path: str) -> str:
    """Convert PDF to DOCX with text extraction and OCR fallback."""
    try:
        from .formatted import convert_formatted_pdf_to_docx
        import pdfplumber  # noqa: F401 - optional layout backend availability
        import docx  # noqa: F401
    except ImportError:
        pass
    else:
        try:
            return convert_formatted_pdf_to_docx(input_path, output_path)
        except Exception:
            logger.warning("Konversi DOCX berformat gagal; memakai ekstraksi teks", exc_info=True)
    workspace = new_workspace(prefix="pdf-docx-")
    try:
        from .lib.docconv.word import write_docx
        
        paragraphs = []
        for index, (text, _tables) in enumerate(_pdf_pages(input_path, workspace)):
            if index:
                paragraphs.append(("", "pagebreak"))
            for line in text.splitlines():
                paragraphs.append(line)
        
        write_docx(paragraphs, output_path)
        return output_path
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def convert_pdf_to_xlsx(input_path: str, output_path: str) -> str:
    """Convert PDF to XLSX with table detection, formatting, and OCR fallback."""
    try:
        from .formatted_xlsx import convert_formatted_pdf_to_xlsx
        import pdfplumber  # noqa: F401 - optional formatting backend
        import openpyxl  # noqa: F401
    except ImportError:
        pass
    else:
        try:
            return convert_formatted_pdf_to_xlsx(input_path, output_path)
        except Exception:
            logger.warning("Konversi XLSX berformat gagal; memakai ekstraksi teks", exc_info=True)
    
    workspace = new_workspace(prefix="pdf-xlsx-")
    try:
        from .lib.docconv.excel import Sheet, write_xlsx
        
        sheets = []
        for index, (text, tables) in enumerate(_pdf_pages(input_path, workspace), 1):
            sheet = Sheet(f"Page {index}")
            if tables:
                row_index = 1
                for table in tables:
                    for row in table:
                        for column, value in enumerate(row):
                            if value is not None and str(value).strip():
                                sheet.set(row_index, column, value)
                        row_index += 1
                    row_index += 1  # blank separator
            else:
                for row_index, line in enumerate(text.splitlines(), 1):
                    if line.strip():
                        sheet.set(row_index, 0, line)
            sheets.append(sheet)
        
        write_xlsx(sheets, output_path)
        return output_path
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def convert_pdf_to_pptx(input_path: str, output_path: str) -> str:
    """Convert PDF to PPTX with editable content, tables, and formatting."""
    try:
        from .formatted_pptx import convert_formatted_pdf_to_pptx
        import pdfplumber  # noqa: F401 - optional formatting backend
        import pptx  # noqa: F401
    except ImportError:
        pass
    else:
        try:
            return convert_formatted_pdf_to_pptx(input_path, output_path)
        except Exception:
            logger.warning("Konversi PPTX berformat gagal; memakai render image", exc_info=True)

    workspace = new_workspace(prefix="pdf-pptx-")
    try:
        from .lib.docconv.imgcodec import read_dimensions_file
        from .lib.docconv.slides import Slide, write_pptx
        
        paths = _render_pdf(input_path, workspace, dpi=150)
        slide_width = 9144000
        first_w, first_h = read_dimensions_file(paths[0])
        slide_height = max(914400, min(51206400,
                                       round(slide_width * first_h / first_w)))
        
        slides, pictures = [], []
        for path in paths:
            width, height = read_dimensions_file(path)
            scale = min(slide_width / width, slide_height / height)
            w, h = max(1, round(width * scale)), max(1, round(height * scale))
            left, top = (slide_width - w) // 2, (slide_height - h) // 2
            slide = Slide()
            slide.add_picture(left, top, w, h)
            slides.append(slide)
            pictures.append([path])
        
        write_pptx(slides, output_path, slide_width, slide_height, pictures)
        return output_path
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def convert_pdf_to_images(input_path: str, output_dir: str) -> list[str]:
    """Render PDF to PNG images."""
    return _render_pdf(input_path, output_dir, dpi=200)


def convert_image_to_pdf(input_path: str, output_path: str) -> str:
    """Convert image (JPG/PNG) to PDF."""
    from .lib.docconv.imgcodec import ImageError, image_to_pdf
    
    try:
        return image_to_pdf(input_path, output_path, dpi=100.0)
    except ImageError as exc:
        raise RuntimeError(str(exc)) from exc


def _office_text_pdf(input_path: str, output_path: str) -> str:
    """Fallback: convert Office files to text-only PDF."""
    from .lib.docconv.pdfwrite import write_text_pdf
    
    story = []
    
    def add(text):
        if str(text).strip():
            story.append(str(text))
    
    def break_page(text=""):
        story.append((text, True))
    
    ext = Path(input_path).suffix.lower()
    if ext == ".docx":
        from .lib.docconv.word import read_docx
        with read_docx(input_path) as reader:
            for kind, value in reader.iter_blocks():
                if kind == "paragraph":
                    add(value)
                else:
                    for row in value:
                        add(" | ".join(cell for cell in row))
    elif ext == ".xlsx":
        from .lib.docconv.excel import read_xlsx
        with read_xlsx(input_path) as reader:
            for index, (title, rows) in enumerate(reader.sheets()):
                if index:
                    break_page()
                add(title)
                for row in rows:
                    if any(value for value in row):
                        add(" | ".join(value for value in row))
    elif ext == ".pptx":
        from .lib.docconv.slides import read_pptx
        with read_pptx(input_path) as reader:
            for index, slide in enumerate(reader.slides()):
                if index:
                    break_page()
                for text in slide.texts:
                    add(text)
    else:
        raise RuntimeError(f"Format Office {ext} belum didukung.")
    
    if not any(isinstance(item, str) and item.strip() for item in story):
        raise RuntimeError("Dokumen tidak memiliki teks. Pasang LibreOffice untuk merender layout.")
    
    write_text_pdf(story, output_path)
    return output_path


def convert_office_to_pdf(input_path: str, output_dir: str) -> str:
    """Convert Office file (DOCX/XLSX/PPTX) to PDF via LibreOffice, enhanced fallback, or text fallback."""
    input_path = str(Path(input_path).resolve())
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    output_path = str(Path(output_dir) / (Path(input_path).stem + ".pdf"))
    
    lo_bin = get_libreoffice_bin()
    if lo_bin:
        workspace = new_workspace(prefix="office-pdf-")
        try:
            profile = (Path(workspace) / "profile").as_uri()
            cmd = [lo_bin, f"-env:UserInstallation={profile}", "--headless",
                   "--convert-to", "pdf", "--outdir", workspace, input_path]
            try:
                subprocess.run(cmd, check=True, capture_output=True, timeout=120)
                result = Path(workspace) / (Path(input_path).stem + ".pdf")
                if result.is_file() and result.stat().st_size:
                    shutil.copyfile(result, output_path)
                    return output_path
                logger.warning("LibreOffice tidak menghasilkan PDF; menggunakan fallback enhanced")
            except (OSError, subprocess.SubprocessError):
                logger.warning("LibreOffice gagal; menggunakan fallback enhanced", exc_info=True)
        finally:
            shutil.rmtree(workspace, ignore_errors=True)
    
    # Try enhanced fallback first
    try:
        from .enhanced_office_pdf import convert_office_to_enhanced_pdf
        import reportlab  # noqa: F401
        return convert_office_to_enhanced_pdf(input_path, output_path)
    except Exception:
        logger.warning("Enhanced fallback gagal; menggunakan text fallback", exc_info=True)
    
    # Final fallback: text-only
    return _office_text_pdf(input_path, output_path)
